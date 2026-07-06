# fosslight xlsx 리포트를 GUI용 gui_result.json으로 정규화
# 시트/컬럼 구조는 RECON.md 참고 (fosslight-scanner v2.1.25 실측 기준)
import glob
import json
import os
from datetime import datetime

from openpyxl import load_workbook

# 시트명 → gui_result.json의 스캐너 키
SHEET_TO_SCANNER = {
    "SRC_FL_Source": "source",
    "BIN_FL_Binary": "binary",
    "DEP_FL_Dependency": "dependency",
}

# xlsx 컬럼명 → OssItem 필드
COLUMN_MAP = {
    "OSS Name": "name",
    "OSS Version": "version",
    "License": "license",
    "Download Location": "downloadLocation",
    "Homepage": "homepage",
    "Copyright Text": "copyright",
    "Exclude": "exclude",
    "Comment": "comment",
}
PATH_COLUMNS = ("Source Path", "Binary Path", "Package URL")


def _parse_sheet(ws):
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    idx = {h: i for i, h in enumerate(header) if h}

    items = []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        # NVD 취약점 검색 링크는 렌더러가 name/version으로 생성 (NVD URL 변경에 유연)
        item = {
            "name": "", "version": "", "license": [], "downloadLocation": "",
            "homepage": "", "copyright": "", "exclude": False, "comment": "",
            "paths": [],
        }
        for col, field in COLUMN_MAP.items():
            if col not in idx:
                continue
            value = row[idx[col]]
            text = "" if value is None else str(value).strip()
            if field == "license":
                item["license"] = [x.strip() for x in text.split(",") if x.strip()]
            elif field == "exclude":
                item["exclude"] = bool(text)
            else:
                item[field] = text
        for col in PATH_COLUMNS:
            if col in idx and row[idx[col]]:
                item["paths"].append(str(row[idx[col]]).strip())
        items.append(item)
    return items


def normalize_report(output_dir, analyzed_path, mode_list):
    tool_info = {}
    items = {"source": [], "binary": [], "dependency": []}

    # 검출 항목이 0건이면 fosslight가 최종 리포트 파일을 생성하지 않음 → 빈 결과로 처리
    reports = glob.glob(os.path.join(output_dir, "fosslight_report_*.xlsx"))
    if reports:
        latest = max(reports, key=os.path.getmtime)
        wb = load_workbook(latest, read_only=True)
        for ws in wb.worksheets:
            if ws.title in SHEET_TO_SCANNER:
                items[SHEET_TO_SCANNER[ws.title]] = _parse_sheet(ws)
            elif ws.title == "Scanner Info":
                for row in ws.iter_rows(values_only=True):
                    if row and row[0] and row[1] is not None:
                        tool_info[str(row[0])] = str(row[1])
        wb.close()

    result = {
        "scanDate": datetime.now().isoformat(timespec="seconds"),
        "analyzedPath": analyzed_path,
        "modes": mode_list,
        "toolInfo": tool_info,
        "items": items,
    }
    result_path = os.path.join(output_dir, "gui_result.json")
    with open(result_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    return result_path
